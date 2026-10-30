"""
Exportación de reportes a Excel (HU25). Es presentación, igual que los
serializers: recibe el DTO que ya armó el caso de uso (con sus permisos y
reglas aplicados) y lo escribe como .xlsx. No calcula nada nuevo, así el
Excel siempre muestra lo mismo que la pantalla.
"""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable, Optional, Sequence

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..domain.entities import ClasificacionABC, CierreDiario, ReporteAntiguedad

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Mismos textos que la pantalla (frontend/src/features/movimientos/formato.ts).
NOMBRE_TIPO = {
    "ENTRADA": "Entrada",
    "SALIDA": "Salida",
    "AJUSTE_POSITIVO": "Ajuste +",
    "AJUSTE_NEGATIVO": "Ajuste −",
    "TRANSFERENCIA_SALIDA": "Transferencia enviada",
    "TRANSFERENCIA_ENTRADA": "Transferencia recibida",
}

FORMATO_MONTO = "#,##0.00"
FORMATO_PORCENTAJE = '0.00"%"'
FORMATO_FECHA_HORA = "dd/mm/yyyy hh:mm"


@dataclass(frozen=True)
class Columna:
    titulo: str
    ancho: int = 14
    formato: Optional[str] = None


def _a_celda(valor: Any) -> Any:
    # Excel no guarda zona horaria: las fechas van en hora de La Paz, como
    # en la pantalla.
    if isinstance(valor, datetime):
        return timezone.localtime(valor).replace(tzinfo=None)
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def _libro(
    titulo: str,
    detalle: str,
    columnas: Sequence[Columna],
    filas: Iterable[Sequence[Any]],
    total: Optional[Sequence[Any]] = None,
) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]  # límite de Excel para el nombre de la hoja
    hoja.append([titulo])
    hoja["A1"].font = Font(size=14, bold=True)
    hoja.append([detalle])
    hoja["A2"].font = Font(color="6E6E73")
    hoja.append([])

    hoja.append([c.titulo for c in columnas])
    fila_cabecera = hoja.max_row
    for celda in hoja[fila_cabecera]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="F2F2F5")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = f"A{fila_cabecera + 1}"

    for fila in filas:
        hoja.append([_a_celda(v) for v in fila])
    if total is not None:
        hoja.append([_a_celda(v) for v in total])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)

    for indice, columna in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = columna.ancho
        if columna.formato:
            for (celda,) in hoja.iter_rows(min_row=fila_cabecera + 1, min_col=indice, max_col=indice):
                celda.number_format = columna.formato

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def respuesta_excel(contenido: bytes, nombre: str) -> HttpResponse:
    respuesta = HttpResponse(contenido, content_type=TIPO_XLSX)
    respuesta["Content-Disposition"] = f'attachment; filename="{slugify(nombre)}.xlsx"'
    return respuesta


def _generado() -> str:
    return f"Generado el {timezone.localtime().strftime('%d/%m/%Y %H:%M')} (hora de La Paz)"


def excel_cierre_diario(cierre: CierreDiario) -> tuple[bytes, str]:
    columnas = [
        Columna("Hora", 18, FORMATO_FECHA_HORA), Columna("Tipo", 22), Columna("Código", 14),
        Columna("Herramienta", 36), Columna("Cantidad", 10), Columna("Unidad", 10),
        Columna("Unidades", 10), Columna("Usuario", 16), Columna("Motivo", 36),
    ]
    filas = [
        (l.fecha_hora, NOMBRE_TIPO.get(l.tipo_movimiento, l.tipo_movimiento), l.codigo, l.nombre, l.cantidad,
         "Caja" if l.tipo_unidad == "CAJA" else "Unidad", l.cantidad_unidades, l.usuario, l.motivo)
        for l in cierre.lineas
    ]
    total = ("Total de salidas del día", "", "", "", "", "", cierre.total_salidas_unidades, "", "")
    detalle = f"{cierre.sucursal_nombre} · {cierre.fecha:%d/%m/%Y} · {_generado()}"
    nombre = f"cierre-diario-{cierre.sucursal_nombre}-{cierre.fecha.isoformat()}"
    return _libro("Cierre diario", detalle, columnas, filas, total), nombre


def excel_antiguedad(reporte: ReporteAntiguedad, sucursal_nombre: Optional[str]) -> tuple[bytes, str]:
    columnas = [
        Columna("Código", 14), Columna("Herramienta", 36), Columna("Sucursal", 16), Columna("Stock", 10),
        Columna("Última venta", 18, FORMATO_FECHA_HORA), Columna("Días sin venta", 14), Columna("Tramo", 16),
        Columna("Estancada", 11), Columna("Costo unitario", 14, FORMATO_MONTO),
        Columna("Valor inmovilizado", 18, FORMATO_MONTO),
    ]
    filas = [
        (l.codigo, l.nombre, l.sucursal_nombre, l.stock, l.ultima_venta or "Nunca vendida", l.dias_sin_venta,
         l.tramo, "Sí" if l.estancada else "No", l.costo_unitario, l.valor_inmovilizado)
        for l in reporte.lineas
    ]
    total = ("Total", "", "", reporte.total_unidades, "", "", "", reporte.total_estancadas, "",
             reporte.valor_inmovilizado_estancado)
    alcance = sucursal_nombre or "Todas las sucursales"
    filtro = f" · desde {reporte.dias_minimos} días sin venta" if reporte.dias_minimos else ""
    detalle = (f"{alcance}{filtro} · estancada = más de {reporte.dias_para_estancamiento} días; el total de "
               f"valor es solo de las estancadas · {_generado()}")
    nombre = f"antiguedad-{alcance}-{reporte.fecha.isoformat()}"
    return _libro("Antigüedad del stock", detalle, columnas, filas, total), nombre


def excel_abc(clasificacion: Sequence[ClasificacionABC], hoy: date) -> tuple[bytes, str]:
    columnas = [
        Columna("Clase", 8), Columna("Código", 14), Columna("Herramienta", 36),
        Columna("Unidades vendidas", 18), Columna("% del total", 12, FORMATO_PORCENTAJE),
        Columna("% acumulado", 12, FORMATO_PORCENTAJE),
    ]
    filas = [
        (c.clase, c.codigo, c.nombre, c.unidades_vendidas, c.porcentaje, c.porcentaje_acumulado)
        for c in clasificacion
    ]
    total = ("Total", "", "", sum(c.unidades_vendidas for c in clasificacion), "", "")
    detalle = f"Por unidades vendidas en los últimos 12 meses · A < 80 %, B < 95 % · {_generado()}"
    return _libro("Clasificación ABC", detalle, columnas, filas, total), f"clasificacion-abc-{hoy.isoformat()}"
